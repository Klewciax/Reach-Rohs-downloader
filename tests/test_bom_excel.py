"""Realistyczne BoM-y z Excela: wiele kart, pusta okładka, nietypowe nagłówki, skróty MPN."""
import re

import pytest
import responses
from openpyxl import Workbook

from bom_compliance.bom import group_items, read_bom
from bom_compliance.config import DEFAULT_MANUFACTURERS, Settings
from bom_compliance.manufacturers import ManufacturerRegistry
from bom_compliance.models import BomItem, Scope
from bom_compliance.mpn import (FORM_ABBREV, FORM_FULL, clean_mpn, hints_from_text, packaging_trims,
                                wildcard_regex)
from bom_compliance.pipeline import Pipeline
from bom_compliance.http_client import PoliteSession
from conftest import make_pdf

REG = ManufacturerRegistry.from_yaml(DEFAULT_MANUFACTURERS)


def _wb(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def test_empty_cover_then_bom_on_second_sheet_polish_headers(tmp_path):
    p = _wb(tmp_path / "b.xlsx", [
        ("Okładka", [[None, "Projekt X"], [None, "Rev C"]]),
        ("BOM", [["Lp", "Oznaczenie", "Opis", "Producent", "Nr kat.", "Ilość"],
                 [1, "U1", "OPAMP LM358DR SOIC8", "TI", "LM358", 1],
                 [2, "Q1, Q2", "MOSFET", "ON Semiconductor", "NDS331N", 2],
                 [3, "L1", "Dławik", "Würth", 7443556082, 1]]),
        ("Historia", [["Rev", "Data"], ["A", "2024-01-01"]]),
    ])
    rows, invalid, meta = read_bom(p, registry=REG)
    assert meta["sheet"] == "BOM" and meta["columns"]["mpn"] == "Nr kat."
    assert any("Okładka" in s and "WYBRANY" not in s for s in meta["sheets"])
    items = {i.mpn: i for i in group_items(rows, REG)}
    assert items["LM358"].hints == ["LM358DR"]                     # pełny numer znaleziony w opisie
    assert items["7443556082"].manufacturer.key == "wurth"
    assert any("liczba" in n for n in items["7443556082"].mpn_notes)  # MPN zapisany jako liczba
    assert items["NDS331N"].refdes == ["Q1", "Q2"]


def test_largest_sheet_chosen_and_sheet_all_merges(tmp_path):
    p = _wb(tmp_path / "b.xlsx", [
        ("Info", []),
        ("Main board", [["Ref", "Manufacturer", "MPN"], ["U1", "TI", "LM358DR"], ["U2", "TI", "TLV70033DDCR"],
                        ["Q1", "onsemi", "NDS331N"]]),
        ("Display board", [["Ref", "Manufacturer", "MPN"], ["U1", "TI", "LM358DR"]]),
    ])
    _, _, meta = read_bom(p, registry=REG)
    assert meta["sheet"] == "Main board"
    rows, _, meta = read_bom(p, sheet="all", registry=REG)
    items = group_items(rows, REG)
    lm = next(i for i in items if i.mpn == "LM358DR")
    assert len(items) == 3 and sorted(lm.rows) == ["Display board!2", "Main board!2"]
    rows, _, meta = read_bom(p, sheet="3", registry=REG)
    assert meta["sheet"] == "Display board"
    with pytest.raises(ValueError):
        read_bom(p, sheet="Nie ma", registry=REG)


def test_columns_detected_by_content_when_headers_unknown(tmp_path):
    p = _wb(tmp_path / "b.xlsx", [("Arkusz1", [
        ["Kol A", "Kol B", "Kol C", "Kol D"],
        ["R1, R2", "Rezystor 10k", "Vishay Dale", "CRCW060310K0FKEA"],
        ["U1", "Wzmacniacz", "Texas Instruments", "LM358DR"],
        ["Q1", "Tranzystor", "Nexperia", "BC847B,215"],
        ["C1", "Kondensator", "Murata", "GRM188R71H104KA93D"],
    ])])
    rows, _, meta = read_bom(p, registry=REG)
    assert meta["columns"]["mpn"] == "Kol D" and meta["columns"]["manufacturer"] == "Kol C"
    assert "zawartości" in meta["detection"] and meta["warnings"]
    assert {r.mpn for r in rows} >= {"LM358DR", "BC847B,215"}


def test_header_pointing_at_description_is_corrected_by_content(tmp_path):
    p = _wb(tmp_path / "b.xlsx", [("BOM", [
        ["Ref", "Manufacturer", "Part", "Code"],
        ["U1", "TI", "Op amp dual SOIC", "LM358DR"],
        ["U2", "TI", "LDO 3.3 V regulator", "TLV70033DDCR"],
        ["Q1", "onsemi", "N channel MOSFET", "NDS331N"],
    ])])
    rows, _, meta = read_bom(p, registry=REG)
    assert meta["columns"]["mpn"] == "Code" and meta["warnings"]


def test_two_row_header_and_alternate_sources(tmp_path):
    p = _wb(tmp_path / "b.xlsx", [("BOM", [
        ["Ref", "Manufacturer", "Manufacturer", "Manufacturer 2", "MPN 2"],
        [None, "Name", "Part Number", None, None],
        ["U1", "TI", "LM358DR", "STMicroelectronics", "LM358DT"],
        ["Q1", "onsemi", "NDS331N", None, None],
    ])])
    rows, _, meta = read_bom(p, registry=REG)
    assert meta["header_row"] == 2
    items = {i.mpn: i for i in group_items(rows, REG)}
    assert items["LM358DT"].alternate and items["LM358DT"].manufacturer.key == "stmicroelectronics"
    assert not items["LM358DR"].alternate and meta["alternates"] == 1
    rows, _, meta = read_bom(p, registry=REG, include_alternates=False)
    assert meta["alternates"] == 0


def test_xls_legacy_format(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    book = xlwt.Workbook()
    book.add_sheet("Cover")
    sh = book.add_sheet("BOM")
    for r, row in enumerate([["Designator", "Mfr", "Mfr Part Number"], ["U1", "TI", "LM358DR"]]):
        for c, v in enumerate(row):
            sh.write(r, c, v)
    p = tmp_path / "b.xls"
    book.save(str(p))
    rows, _, meta = read_bom(p, registry=REG)
    assert meta["sheet"] == "BOM" and rows[0].mpn == "LM358DR"


def test_clean_mpn_cases():
    c = clean_mpn("TI LM358DR", ["Texas Instruments", "TI"])
    assert c.mpn == "LM358DR" and c.notes
    assert clean_mpn("LM358DR (SOIC-8)").mpn == "LM358DR"
    assert clean_mpn("LM358DR - dual opamp").mpn == "LM358DR"
    c = clean_mpn("LM358DR\nLM358DT")
    assert c.mpn == "LM358DR" and c.alternates == ["LM358DT"]
    assert clean_mpn("TJA1051T/3").mpn == "TJA1051T/3"      # "/" w MPN nie dzieli numerów
    assert clean_mpn("BAS16,215").mpn == "BAS16,215"
    c = clean_mpn("CRCW0603xxxxFKEA")
    assert c.wildcard and wildcard_regex(c.mpn).match("CRCW060310K0FKEA")
    assert clean_mpn("STM32F103C8T*").wildcard
    assert hints_from_text("LM358", ["OPAMP LM358DR SOIC8", "foo"]) == ["LM358DR"]
    assert [t for t, _ in packaging_trims("LTC3780EG#PBF")] == ["LTC3780EG"]
    assert [t for t, _ in packaging_trims("BAS16,215")] == ["BAS16"]
    assert packaging_trims("LM358DR") == []


# ------------------------------------------------------------- online (mock)

def _settings(tmp_path):
    return Settings.load(None, {"output_dir": str(tmp_path), "min_delay_per_host": 0, "delay_jitter": 0,
                                "max_retries": 0, "download_general_statements": False, "generic_max_pages": 0})


def _item(mpn, key, hints=()):
    m = REG.manufacturers[key]
    return BomItem(mpn=mpn, manufacturer_raw=m.name, manufacturer=m, match_method="exact", mpn_bom=mpn,
                   hints=list(hints))


@responses.activate
def test_abbreviated_mpn_ambiguous_is_flagged_and_docs_capped_to_family(tmp_path):
    responses.get("https://www.microchip.com/robots.txt", status=404)
    responses.get("https://ww1.microchip.com/robots.txt", status=404)
    responses.get("https://www.microchip.com/en-us/product/ATMEGA328P",
                  body="<html><h1>ATmega328P</h1><table><tr><td>ATMEGA328P-AU</td></tr>"
                       "<tr><td>ATMEGA328P-PU</td></tr><tr><td>ATMEGA328P-MU</td></tr></table></html>",
                  content_type="text/html")
    responses.get("https://ww1.microchip.com/downloads/aemDocuments/documents/corporate-responsibilty/environmental/"
                  "material-compliance-documents/EU-RoHS-CoC_ATMEGA328P.pdf",
                  body=make_pdf("CERTIFICATE OF RoHS COMPLIANCE\nATMEGA328P\nDirective 2011/65/EU"),
                  content_type="application/pdf")
    responses.get(re.compile(r"https://.*"), status=404)
    s = _settings(tmp_path)
    res = Pipeline(s, tmp_path, PoliteSession(s, sleep=lambda _: None)).process_item(_item("ATMEGA328P", "microchip"))
    chk = res.item.mpn_check
    assert chk.form == FORM_ABBREV and not chk.expanded_to
    assert set(chk.variants) == {"ATMEGA328P-AU", "ATMEGA328P-PU", "ATMEGA328P-MU"}
    assert res.docs and res.docs[0].scope == Scope.FAMILY  # dokument nie potwierdza konkretnego wariantu


@responses.activate
def test_abbreviated_mpn_expanded_with_hint_from_description(tmp_path):
    responses.get("https://www.microchip.com/robots.txt", status=404)
    responses.get("https://www.microchip.com/en-us/product/ATMEGA328P",
                  body="<html><h1>ATmega328P</h1>ATMEGA328P-AU ATMEGA328P-PU</html>", content_type="text/html")
    responses.get(re.compile(r"https://.*"), status=404)
    s = _settings(tmp_path)
    item = _item("ATMEGA328P", "microchip", hints=["ATMEGA328P-AU"])
    res = Pipeline(s, tmp_path, PoliteSession(s, sleep=lambda _: None)).process_item(item)
    assert res.item.mpn_check.expanded_to == "ATMEGA328P-AU" and res.item.mpn == "ATMEGA328P-AU"
    assert res.item.mpn_bom == "ATMEGA328P"


@responses.activate
def test_full_mpn_confirmed_on_orderable_page(tmp_path):
    responses.get("https://www.ti.com/robots.txt", status=404)
    responses.get("https://www.ti.com/store/ti/en/p/product/?p=LM358DR",
                  body="<html>LM358DR Active LM358DRG3 Obsolete</html>", content_type="text/html")
    responses.get(re.compile(r"https://.*"), status=404)
    s = _settings(tmp_path)
    res = Pipeline(s, tmp_path, PoliteSession(s, sleep=lambda _: None)).process_item(_item("LM358DR", "texas_instruments"))
    assert res.item.mpn_check.form == FORM_FULL and res.item.mpn == "LM358DR"


@responses.activate
def test_packaging_suffix_fallback_marks_family(tmp_path):
    responses.get("https://www.onsemi.com/robots.txt", status=404)
    responses.get("https://onsemi.com/robots.txt", status=404)
    responses.get("https://www.onsemi.com/PowerSolutions/MaterialComposition.do?export=coc_pdf&opnId=NDS331N",
                  body=make_pdf("CERTIFICATE OF COMPLIANCE RoHS and REACH\nNDS331N\nSVHC"),
                  content_type="application/pdf")
    responses.get(re.compile(r"https://.*"), status=404)
    s = _settings(tmp_path)
    res = Pipeline(s, tmp_path, PoliteSession(s, sleep=lambda _: None)).process_item(_item("NDS331N-TR", "onsemi"))
    assert res.docs and res.docs[0].scope == Scope.FAMILY
    assert any("NDS331N" in n and "sufiks" in n.lower() or "formy NDS331N" in n for n in res.notes)
