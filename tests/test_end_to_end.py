"""Test end-to-end z zamockowanymi stronami producentów (bez dostępu do sieci)."""
import csv
import json
import re
from pathlib import Path

import responses
from openpyxl import load_workbook

from bom_compliance.cli import main
from conftest import make_pdf

ONSEMI_COC = "https://www.onsemi.com/PowerSolutions/MaterialComposition.do?export=coc_pdf&opnId=NDS331N"
ONSEMI_MCD = "https://www.onsemi.com/PowerSolutions/MaterialComposition.do?export=pdf&opnId=NDS331N"


def _mock_sites():
    for host in ("www.onsemi.com", "www.nexperia.com", "www.vishay.com", "www.ti.com"):
        responses.get(f"https://{host}/robots.txt", body="User-agent: *\nDisallow: /private/\n")

    # onsemi – stały wzorzec URL (API eksportu)
    responses.get(ONSEMI_COC, body=make_pdf("CERTIFICATE OF COMPLIANCE RoHS and REACH\nPart Number: NDS331N\n"
                                            "Compliant with Directive 2011/65/EU and 2015/863\nNo SVHC > 0.1%"),
                  content_type="application/pdf")
    responses.get(ONSEMI_MCD, body=make_pdf("Material Composition Declaration\nNDS331N\nIPC-1752A"),
                  content_type="application/pdf")

    # Nexperia – numer zamówieniowy nie ma strony, numer typu ma
    responses.get("https://www.nexperia.com/chemical-content/BAS16%2C215.html", status=404)
    responses.get("https://www.nexperia.com/chemical-content/BAS16.html",
                  body='<html><body><h1>Chemical content for BAS16</h1>'
                       '<a href="/docs/chemical-content/BAS16_chemical_content.pdf">Download PDF</a>'
                       '<a href="https://www.digikey.com/rohs/BAS16.pdf">RoHS at distributor</a></body></html>',
                  content_type="text/html")
    responses.get("https://www.nexperia.com/docs/chemical-content/BAS16_chemical_content.pdf",
                  body=make_pdf("Chemical content BAS16\nRoHS compliant\nREACH SVHC: none"),
                  content_type="application/pdf")

    # Vishay – generyczny crawl strony compliance
    responses.get("https://www.vishay.com/en/how/leadfree/",
                  body='<html><body>'
                       '<a href="/docs/28429/crcwrohs.pdf">RoHS Certificate thick film chip resistors</a>'
                       '<a href="https://www.mouser.com/datasheet/vishay-rohs.pdf">RoHS (Mouser)</a>'
                       '<a href="/en/quality/">Quality</a></body></html>',
                  content_type="text/html")
    responses.get("https://www.vishay.com/docs/28429/crcwrohs.pdf",
                  body=make_pdf("RoHS Certificate\nVishay Dale CRCW0603 e3 series\nDirective 2011/65/EU"),
                  content_type="application/pdf")
    responses.get("https://www.vishay.com/en/quality/",
                  body='<html><body><a href="mailto:env-compliance@vishay.com">Environmental compliance</a>'
                       ' contact also fake@gmail.com</body></html>',
                  content_type="text/html")

    # TI – per MPN tylko po zalogowaniu; ogólne oświadczenia publiczne
    responses.get("https://www.ti.com/lit/pdf/szzq087",
                  body=make_pdf("Statement on REACH Articles Provisions\nTexas Instruments\nSVHC"),
                  content_type="application/pdf")
    responses.get("https://www.ti.com/lit/pdf/szzq088",
                  body=make_pdf("TI RoHS and Product Statement Designation\nRoHS 2011/65/EU"),
                  content_type="application/pdf")
    for url in ("https://www.ti.com/quality-reliability/environmental-information.html",
                "https://www.ti.com/support-quality/faqs/environmental-information-faqs.html",
                "https://www.ti.com/support.html"):
        responses.get(url, body="<html><body>Environmental info</body></html>", content_type="text/html")
    responses.get("https://www.ti.com/materialcontent/home", status=302,
                  headers={"Location": "https://login.ti.com/as/authorization.oauth2"})
    responses.get(re.compile(r"https://.*"), status=404)  # wszystko inne


@responses.activate
def test_full_run(tmp_path):
    _mock_sites()
    cfg = tmp_path / "cfg.yaml"
    out = tmp_path / "out"
    cfg.write_text("min_delay_per_host: 0\ndelay_jitter: 0\nbackoff_base: 0\nmax_retries: 1\n")
    bom = tmp_path / "bom.csv"
    bom.write_text(
        "RefDes,Mfr,MPN\n"
        "\"Q1,Q2\",ON Semiconductor,NDS331N\n"
        "Q3,onsemi,NDS331N\n"
        "D1,Nexperia,\"BAS16,215\"\n"
        "R1,Vishay Dale,CRCW060310K0FKEA\n"
        "U1,Texas Instruments Inc.,LM358DR\n"
        "U2,TI,TLV70033DDCR\n"
        "J1,Acme Connectors Ltd,AC-1234\n"
        "R2,,RC0603FR-0710KL\n",
        encoding="utf-8")

    rc = main([str(bom), "-o", str(out), "-c", str(cfg)])
    assert rc == 0

    # Nigdy nie odpytano domen spoza producentów
    hosts = {re.match(r"https?://([^/]+)", c.request.url).group(1) for c in responses.calls}
    assert hosts <= {"www.onsemi.com", "www.nexperia.com", "www.vishay.com", "www.ti.com"}, hosts
    # Deduplikacja: certyfikat onsemi pobrany raz mimo 3 refdes w 2 wierszach
    assert sum(1 for c in responses.calls if c.request.url == ONSEMI_COC) == 1
    # Ogólne oświadczenie TI pobrane raz dla dwóch MPN
    assert sum(1 for c in responses.calls if c.request.url.endswith("szzq087")) == 1

    rows = {r["MPN"]: r for r in csv.DictReader(open(out / "report_items.csv", encoding="utf-8-sig"), delimiter=";")}
    assert rows["NDS331N"]["Status RoHS"].startswith("POBRANO (")
    assert rows["NDS331N"]["Status REACH"].startswith("POBRANO (")
    assert rows["NDS331N"]["URL źródłowy RoHS"] == ONSEMI_COC
    assert rows["NDS331N"]["Ref. designators"] == "Q1, Q2, Q3"
    assert rows["BAS16,215"]["Status RoHS"].startswith("POBRANO – DOKUMENT ZBIORCZY")
    assert rows["CRCW060310K0FKEA"]["Status RoHS"].startswith("POBRANO – DOKUMENT ZBIORCZY")
    assert rows["CRCW060310K0FKEA"]["Status REACH"] == "NIE ZNALEZIONO"
    assert "mouser" not in rows["CRCW060310K0FKEA"]["URL źródłowy RoHS"]
    assert rows["LM358DR"]["Status RoHS"] == "WYMAGA LOGOWANIA"
    assert rows["AC-1234"]["Status RoHS"] == "NIEZNANY PRODUCENT"

    # Plik + metadane źródła
    pdf = Path(rows["NDS331N"]["Plik RoHS"])
    assert pdf.is_file() and pdf.parent.name == "NDS331N" and pdf.name.startswith("NDS331N__onsemi__RoHS-REACH")
    meta = json.loads(pdf.with_name(pdf.name + ".source.json").read_text())
    assert meta["source_url"] == ONSEMI_COC and meta["downloaded_at_utc"]
    ti_general = [p for p in (out / "documents" / "Texas_Instruments" / "_shared").iterdir() if p.suffix == ".pdf"]
    assert len(ti_general) == 2

    # Lista "do uzyskania mailowo": pogrupowana per producent, e-mail tylko z domeny producenta
    req = list(csv.DictReader(open(out / "report_to_request.csv", encoding="utf-8-sig"), delimiter=";"))
    by_mfr = {}
    for r in req:
        by_mfr.setdefault(r["Producent"], []).append(r)
    assert {r["MPN"] for r in by_mfr["Texas Instruments"]} == {"LM358DR", "TLV70033DDCR"}
    assert "NDS331N" not in {r["MPN"] for r in req}
    assert "env-compliance@vishay.com" in by_mfr["Vishay"][0]["Kontakt (oficjalna strona producenta)"]
    assert "gmail" not in by_mfr["Vishay"][0]["Kontakt (oficjalna strona producenta)"]
    assert "DO RĘCZNEJ WERYFIKACJI" in by_mfr["Acme Connectors Ltd"][0]["Kontakt (oficjalna strona producenta)"]
    tpl = (out / "email_templates" / "Texas_Instruments.txt").read_text()
    assert "LM358DR" in tpl and "TLV70033DDCR" in tpl and "SVHC" in tpl

    wb = load_workbook(out / "report.xlsx")
    assert {"Podsumowanie", "Pozycje", "Pliki", "Do uzyskania mailowo", "Szablony e-mail",
            "Niepoprawne wiersze"} <= set(wb.sheetnames)
    summary = {r[0]: r[1:] for r in wb["Podsumowanie"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Pozycje BoM po deduplikacji (producent + MPN)"][0] == "6"
    assert summary["Pobrano oba (RoHS i REACH)"][0] == "2"  # onsemi + Nexperia (rodzina)
    assert wb["Niepoprawne wiersze"].max_row == 2
