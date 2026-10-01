"""Smoke testy: szybkie sprawdzenie, że narzędzie jako całość się uruchamia.

Offline (zawsze):  pytest -m smoke
Live (opcjonalnie, prawdziwe strony producentów):  BOM_LIVE_SMOKE=1 pytest -m live -s
"""
import os
import re
import subprocess
import sys
from pathlib import Path

import pytest
import responses
import yaml
from openpyxl import load_workbook

from bom_compliance import smoke
from bom_compliance.cli import main
from bom_compliance.config import PACKAGE_ROOT, Settings

ROOT = PACKAGE_ROOT
pytestmark = pytest.mark.smoke


def test_offline_smoke_has_no_broken_checks():
    checks = smoke.offline_checks()
    broken = [c for c in checks if c.level == smoke.BROKEN]
    assert not broken, broken
    assert smoke.main([]) == 0


def test_smoke_detects_url_outside_manufacturer_domain(tmp_path):
    reg = tmp_path / "m.yaml"
    reg.write_text(yaml.safe_dump({"manufacturers": {"acme": {
        "name": "ACME", "domains": ["acme.com"], "adapter": "generic",
        "compliance_pages": ["https://www.mouser.com/acme-rohs"],
        "product_pages": ["https://www.acme.com/p/{unknown_field}"]}}}))
    checks = smoke.offline_checks(reg)
    msgs = " ".join(c.message for c in checks if c.level == smoke.BROKEN)
    assert "mouser.com" in msgs and "unknown_field" in msgs


def _run(*args):
    return subprocess.run([sys.executable, "-m", "bom_compliance", *args], cwd=ROOT, capture_output=True,
                          text=True, timeout=60)


def test_cli_help():
    r = _run("--help")
    assert r.returncode == 0 and "--lifecycle" in r.stdout and "--longevity" in r.stdout


def test_cli_dry_run_example_bom():
    r = _run("examples/bom_example.csv", "--dry-run")
    assert r.returncode == 0, r.stderr
    assert "13 unikalnych pozycji" in r.stdout and "NDS331N" in r.stdout


def test_cli_bad_input_returns_error_code(tmp_path):
    assert _run(str(tmp_path / "missing.csv")).returncode == 2
    bad = tmp_path / "bad.csv"
    bad.write_text("a,b\n1,2\n")
    assert _run(str(bad)).returncode == 2


@responses.activate
def test_example_bom_full_run_all_options_offline(tmp_path):
    """Cały przebieg na przykładowym BoM, gdy żadna strona producenta nie odpowiada sensownie (404)."""
    responses.get(re.compile(r"https?://.*"), status=404)
    cfg = tmp_path / "cfg.yaml"
    cfg.write_text("min_delay_per_host: 0\ndelay_jitter: 0\nbackoff_base: 0\nmax_retries: 0\n")
    out = tmp_path / "out"
    rc = main([str(ROOT / "examples" / "bom_example.csv"), "-o", str(out), "-c", str(cfg),
               "--lifecycle", "--longevity"])
    assert rc == 0
    for name in ("report.xlsx", "report_items.csv", "report_files.csv", "report_to_request.csv",
                 "report_lifecycle.csv", "run.log"):
        assert (out / name).is_file(), name
    assert any((out / "email_templates").iterdir())
    sheets = load_workbook(out / "report.xlsx").sheetnames
    assert {"Podsumowanie", "Pozycje", "Pliki", "Do uzyskania mailowo", "Szablony e-mail",
            "Niepoprawne wiersze", "Cykl życia i longevity"} <= set(sheets)
    # Żadne zapytanie nie wyszło poza domeny producentów z rejestru
    reg = yaml.safe_load((ROOT / "config" / "manufacturers.yaml").read_text())["manufacturers"]
    domains = [d for m in reg.values() for d in m["domains"]]
    for call in responses.calls:
        host = re.match(r"https?://([^/]+)", call.request.url).group(1)
        assert any(host == d or host.endswith("." + d) for d in domains), host


@responses.activate
def test_live_smoke_logic_flags_dead_urls(tmp_path):
    """Logika trybu --live (na mockach): martwy URL = BROKEN, kod wyjścia 1."""
    reg = tmp_path / "m.yaml"
    reg.write_text(yaml.safe_dump({"manufacturers": {"acme": {
        "name": "ACME", "domains": ["acme.com"], "adapter": "generic",
        "compliance_pages": ["https://www.acme.com/env", "https://www.acme.com/old-page"]}}}))
    parts = tmp_path / "parts.yaml"
    parts.write_text("parts: {}\n")
    responses.get("https://www.acme.com/robots.txt", status=404)
    responses.get("https://www.acme.com/env", body="<html>ok</html>", content_type="text/html")
    responses.get("https://www.acme.com/old-page", status=404)
    settings = Settings.load(None, {"manufacturers_file": str(reg), "min_delay_per_host": 0, "delay_jitter": 0,
                                    "max_retries": 0})
    live = smoke.live_checks(settings, parts_file=parts)
    levels = {c.message.split()[-1]: c.level for c in live}
    assert levels["https://www.acme.com/env"] == smoke.OK
    assert levels["https://www.acme.com/old-page"] == smoke.BROKEN
    assert smoke.exit_code([], live) == 1


@pytest.mark.live
@pytest.mark.skipif(os.environ.get("BOM_LIVE_SMOKE") != "1", reason="ustaw BOM_LIVE_SMOKE=1, aby odpytać prawdziwe strony")
def test_live_smoke_real_sites():
    only = set(filter(None, os.environ.get("BOM_LIVE_ONLY", "").split(","))) or None
    settings = Settings.load(os.environ.get("BOM_LIVE_CONFIG"), {"check_lifecycle": True, "check_longevity": True,
                                                                 "save_lifecycle_snapshots": False})
    live = smoke.live_checks(settings, only)
    smoke.print_checks(live)
    if not any(c.level == smoke.OK for c in live):
        pytest.skip("brak dostępu do stron producentów z tego środowiska (problem sieci)")
    broken = [c for c in live if c.level == smoke.BROKEN]
    assert not broken, "\n".join(f"{c.area}: {c.message}" for c in broken)
