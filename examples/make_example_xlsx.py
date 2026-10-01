"""Generuje examples/bom_example_multisheet.xlsx – BoM w typowym układzie firmowym:
pusta okładka, właściwa lista na 2. karcie (polskie nagłówki, zamienniki, skróty MPN), historia zmian."""
from pathlib import Path

from openpyxl import Workbook

wb = Workbook()
cover = wb.active
cover.title = "Strona tytułowa"
cover["B2"] = "Urządzenie XYZ – zestawienie materiałowe"
cover["B3"] = "Rewizja C"

bom = wb.create_sheet("BOM")
bom.append(["Zestawienie materiałowe – płyta główna"])
bom.append([])
bom.append(["Lp", "Oznaczenie", "Opis", "Ilość", "Producent", "Nr kat.", "Producent 2", "Nr kat. 2"])
rows = [
    ("U1", "Wzmacniacz operacyjny LM358DR SOIC-8", 1, "Texas Instruments Inc.", "LM358", "STMicroelectronics", "LM358DT"),
    ("U2", "Mikrokontroler", 1, "ST", "STM32F103C8T6", None, None),
    ("Q1, Q2", "MOSFET N", 2, "ON Semiconductor", "NDS331N", None, None),
    ("D1 D2", "Dioda impulsowa", 2, "Nexperia", "BAS16,215", None, None),
    ("U3", "Przetwornica", 1, "Linear Technology", "LTC3780EG#PBF", None, None),
    ("U4", "Mikrokontroler 8-bit", 1, "Microchip", "ATMEGA328P", None, None),
    ("R1-R10", "Rezystor 10k 1% 0603", 10, "Vishay Dale", "CRCW0603xxxxFKEA", "Yageo", "RC0603FR-0710KL"),
    ("C1", "Kondensator 100n", 1, "Murata", "GRM188R71H104KA93D", None, None),
    ("L1", "Dławik", 1, "Würth Elektronik", 7443556082, None, None),
    ("U5", "Transceiver CAN", 1, "NXP", "TJA1051T/3", None, None),
    ("TP1", "Punkt testowy", 1, "Generic", "N/A", None, None),
]
for i, r in enumerate(rows, 1):
    bom.append([i, *r])

hist = wb.create_sheet("Historia zmian")
hist.append(["Rev", "Data", "Opis"])
hist.append(["A", "2025-03-01", "Pierwsze wydanie"])
hist.append(["C", "2026-06-10", "Zmiana U1"])

out = Path(__file__).with_name("bom_example_multisheet.xlsx")
wb.save(out)
print(f"Zapisano {out}")
