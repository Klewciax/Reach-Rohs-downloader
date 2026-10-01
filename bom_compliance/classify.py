"""Klasyfikacja dokumentów: rodzaj (RoHS/REACH/MCD) i zakres (MPN / rodzina / ogólny)."""
from __future__ import annotations

import io
import logging
import re

from .models import DocType, Scope

log = logging.getLogger(__name__)

_ROHS_RE = re.compile(r"\bRoHS\b|2011/65|2015/863", re.IGNORECASE)
# "REACH" wielkimi literami lub numer rozporządzenia / SVHC — unika fałszywych trafień ("reach out").
_REACH_RE = re.compile(r"\bREACH\b|\bREACh\b|1907/2006|\bSVHC\b|Substances? of Very High Concern", re.UNICODE)
_MCD_RE = re.compile(
    r"material\s+(content|composition|declaration)|IPC[\s-]?1752|full\s+material\s+disclosure|chemical\s+content",
    re.IGNORECASE,
)

# Słowa kluczowe w URL / tekście linku (małe litery, granice "nie-liter").
_URL_ROHS = re.compile(r"(?<![a-z])rohs(?![a-z])")
_URL_REACH = re.compile(r"(?<![a-z])(reach|svhc)(?![a-z])")
_URL_MCD = re.compile(
    r"material[-_ ]?(content|composition|declaration)|materialcontent|materialcomposition|"
    r"chemical[-_ ]?content|(?<![a-z])mcds?(?![a-z])|(?<![a-z])md\.(pdf|xml)|ipc[-_ ]?1752|_md(?![a-z])"
)

DOC_EXTENSIONS = (".pdf", ".xml", ".xls", ".xlsx", ".zip", ".doc", ".docx", ".ashx")


def compact(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def types_from_text(text: str) -> set[DocType]:
    out: set[DocType] = set()
    if _ROHS_RE.search(text):
        out.add(DocType.ROHS)
    if _REACH_RE.search(text):
        out.add(DocType.REACH)
    if _MCD_RE.search(text):
        out.add(DocType.MCD)
    return out


def types_from_link(url: str, text: str = "") -> set[DocType]:
    hay = f"{url} {text}".lower()
    out: set[DocType] = set()
    if _URL_ROHS.search(hay):
        out.add(DocType.ROHS)
    if _URL_REACH.search(hay):
        out.add(DocType.REACH)
    if _URL_MCD.search(hay):
        out.add(DocType.MCD)
    return out


def mpn_prefixes(mpn: str, min_len: int = 4) -> list[str]:
    """Coraz krótsze prefiksy MPN (do wykrywania dokumentów rodzinnych)."""
    c = compact(mpn)
    floor = max(min_len, int(len(c) * 0.5))
    out = []
    for n in range(len(c) - 1, floor - 1, -1):
        p = c[:n]
        if re.search(r"[A-Z]", p) and re.search(r"[0-9]", p):
            out.append(p)
    return out


def _mpn_regex(mpn: str) -> re.Pattern | None:
    chars = [re.escape(c) for c in compact(mpn)]
    if not chars:
        return None
    # Separatory (- / . , # spacja) między znakami traktowane elastycznie: "LM358-DR" == "LM358DR".
    return re.compile(r"(?<![A-Za-z0-9])" + r"[-/.,# ]?".join(chars), re.I)


_TOKEN_TAIL = re.compile(r"[-/#+,.]?[A-Za-z0-9][A-Za-z0-9\-/#+,.]*")


def is_exact_end(text: str, end: int) -> bool:
    """Czy dopasowanie MPN kończy się na granicy numeru.

    "LM358DR " / "LM358DR," / "LM358DR." -> koniec numeru;
    "LM358DRG4", "ATMEGA328P-AU", "BAS16,215", "LTC3780EG#PBF" -> dłuższy numer (inny wariant).
    """
    nxt = text[end:end + 1]
    if not nxt:
        return True
    if nxt.isalnum():
        return False
    if nxt in "-/#+,." and text[end + 1:end + 2].isalnum():
        return False
    return True


def find_mpn_tokens(text: str, mpn: str) -> tuple[bool, list[str]]:
    """Szuka MPN w tekście z poszanowaniem granic tokenu.

    Zwraca (exact, longer): exact=True, jeśli MPN występuje jako samodzielny numer;
    longer = dłuższe numery zaczynające się od MPN (np. warianty zamówieniowe LM358 -> LM358DR).
    """
    rx = _mpn_regex(mpn)
    if rx is None:
        return False, []
    exact, longer = False, []
    for m in rx.finditer(text):
        if is_exact_end(text, m.end()):
            exact = True
            continue
        tail = _TOKEN_TAIL.match(text, m.end())
        token = (m.group(0) + (tail.group(0) if tail else "")).rstrip("-/#+,.")
        if len(token) <= len(m.group(0)) + 15 and token.upper() not in (t.upper() for t in longer):
            longer.append(token)
    return exact, longer


def scope_from_text(mpn: str, text: str) -> tuple[Scope, str]:
    """Określa zakres dokumentu na podstawie obecności MPN w treści.

    PART tylko wtedy, gdy MPN występuje jako samodzielny numer. Gdy w treści są wyłącznie
    dłuższe numery (np. BoM: "LM358", dokument: "LM358DRG4"), dokument dotyczy innego
    wariantu i jest oznaczany jako zbiorczy (FAMILY).
    """
    exact, longer = find_mpn_tokens(text, mpn)
    if exact:
        return Scope.PART, f"MPN {mpn} znaleziony w treści"
    if longer:
        return Scope.FAMILY, ("w treści występują tylko dłuższe numery (inne warianty): "
                              + ", ".join(longer[:5]))
    hay = compact(text)
    for p in mpn_prefixes(mpn):
        if p in hay:
            return Scope.FAMILY, f"w treści znaleziono tylko prefiks rodziny '{p}'"
    return Scope.GENERAL, "MPN nie występuje w treści dokumentu"


def scope_from_link(mpn: str, url: str, text: str = "") -> Scope:
    hay = compact(f"{url} {text}")
    if compact(mpn) and compact(mpn) in hay:
        return Scope.PART
    if any(p in hay for p in mpn_prefixes(mpn)):
        return Scope.FAMILY
    return Scope.GENERAL


def sniff_kind(data: bytes, url: str = "", content_type: str = "") -> str:
    head = data[:1024].lstrip()
    ct = content_type.lower()
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if head.startswith(b"PK"):
        low = url.lower().split("?")[0]
        if low.endswith(".docx"):
            return "docx"
        if low.endswith(".zip"):
            return "zip"
        return "xlsx" if (low.endswith(".xlsx") or "spreadsheet" in ct) else "zip"
    lowhead = head[:512].lower()
    if lowhead.startswith(b"<?xml") and b"<html" not in lowhead:
        return "xml"
    if b"<html" in lowhead or b"<!doctype html" in lowhead or "html" in ct:
        return "html"
    return "bin"


def extract_text(data: bytes, kind: str, max_pages: int = 6) -> str:
    try:
        if kind == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            parts = []
            for page in reader.pages[:max_pages]:
                parts.append(page.extract_text() or "")
            return "\n".join(parts)
        if kind in ("xml", "html"):
            text = data.decode("utf-8", errors="ignore")
            if kind == "html":
                from bs4 import BeautifulSoup

                text = BeautifulSoup(text, "html.parser").get_text(" ")
            return text
        if kind == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            cells = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=2000, values_only=True):
                    cells.extend(str(v) for v in row if v is not None)
            return " ".join(cells)
    except Exception as exc:  # uszkodzony / zaszyfrowany plik
        log.warning("Nie udało się wyodrębnić tekstu (%s): %s", kind, exc)
    return ""


def looks_like_login_page(html: str) -> bool:
    low = html.lower()
    return 'type="password"' in low or "type='password'" in low or "type=password" in low
