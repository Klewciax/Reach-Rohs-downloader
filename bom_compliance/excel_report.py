"""Główny raport w Excelu (report.xlsx).

Karty:
  Podsumowanie | RoHS | REACH | Status cyklu życia | Długość produkcji | Szczegóły pozycji |
  Pliki | Do uzyskania mailowo | Szablony e-mail | Niepoprawne wiersze

Na kartach RoHS / REACH / Status cyklu życia / Długość produkcji przy każdym komponencie
(MPN + producent) jest kolumna "Status", której komórka jest linkiem do POBRANEGO PLIKU
w folderze wyjściowym (link względny – działa po przeniesieniu całego folderu raportu).
Adres strony producenta jest podany osobno, jako zwykły tekst (ślad pochodzenia pliku).
"""
from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .downloader import LIFECYCLE_FOLDER, TYPE_FOLDERS
from .models import DocType, ItemResult, LifecycleStatus, Scope, Status

GREEN, YELLOW, ORANGE, RED, GREY = "C6EFCE", "FFF2CC", "FFEB9C", "FFC7CE", "EDEDED"
SCOPE_PL = {Scope.PART: "dla MPN", Scope.FAMILY: "zbiorczy (rodzina / seria)", Scope.GENERAL: "ogólne oświadczenie"}


@dataclass
class Link:
    """Komórka z tekstem i linkiem do pliku lokalnego (ścieżka względem report.xlsx)."""
    text: str
    target: str | None = None
    fill: str | None = None


class _Writer:
    def __init__(self, out_dir: Path):
        self.out_dir = out_dir.resolve()
        self.wb = Workbook()
        self.bold = Font(bold=True)
        self.head_fill = PatternFill("solid", fgColor="DDEBF7")
        self.link_font = Font(color="0563C1", underline="single")

    def rel(self, path: str | None) -> str | None:
        if not path:
            return None
        p = Path(path).resolve()
        try:
            rel = p.relative_to(self.out_dir)
        except ValueError:
            rel = Path(os.path.relpath(p, self.out_dir))
        return rel.as_posix()

    def sheet(self, title: str, header: list[str], rows: list[list], widths: dict[str, int] | None = None,
              first: bool = False):
        ws = self.wb.active if first else self.wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = self.bold, self.head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx)
                if isinstance(value, Link):
                    cell.value = value.text
                    target = self.rel(value.target)
                    if target:
                        cell.hyperlink = target
                        cell.font = self.link_font
                    if value.fill:
                        cell.fill = PatternFill("solid", fgColor=value.fill)
                else:
                    cell.value = value
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2" if len(header) > 3 else "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, name in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(name, 20)
        return ws


# --------------------------------------------------------------- pomocnicze

def _status_fill(status: Status) -> str:
    return {Status.DOWNLOADED: GREEN, Status.DOWNLOADED_FAMILY: YELLOW, Status.GENERAL_ONLY: YELLOW}.get(status, RED)


def _ident(n: int, r: ItemResult) -> list:
    it = r.item
    mpn = it.mpn_bom or it.mpn
    if it.mpn != mpn:
        mpn = f"{mpn} (wyszukano jako {it.mpn})"
    return [n, mpn, it.manufacturer_name, " | ".join(it.manufacturer_variants) or it.manufacturer_raw,
            ", ".join(it.refdes), "TAK" if it.alternate else ""]


IDENT = ["Lp", "MPN", "Producent", "Producent (zapis w BoM)", "Ref. designators", "Zamiennik"]
IDENT_W = {"Lp": 6, "MPN": 26, "Producent": 22, "Producent (zapis w BoM)": 22, "Ref. designators": 18,
           "Zamiennik": 10}


def _compliance_rows(results: list[ItemResult], t: DocType, status_pl: dict) -> list[list]:
    rows = []
    for n, r in enumerate(results, 1):
        status = r.rohs_status if t == DocType.ROHS else r.reach_status
        docs = r.docs_for(t)
        best = [d for d in docs if d.scope == docs[0].scope] if docs else []
        main = best[0] if best else None
        others = [d for d in docs if d is not main]
        notes = list(r.reasons)
        for d in docs:
            if d.note:
                notes.append(f"{Path(d.path_for(t)).name}: {d.note}")
        rows.append([
            *_ident(n, r),
            Link(status_pl[status], main.path_for(t) if main else None, _status_fill(status)),
            SCOPE_PL[main.scope] if main else "",
            Link(Path(main.path_for(t)).name, main.path_for(t)) if main else "",
            "\n".join(Path(d.path_for(t)).name for d in others),
            main.downloaded_at if main else "",
            (main.final_url if main.final_url == main.url else f"{main.url}\n-> {main.final_url}") if main else "",
            main.mpn_verified if main else "",
            "\n".join(dict.fromkeys(notes)),
            "\n".join(r.manual_urls),
        ])
    return rows


COMPLIANCE_HEADER = IDENT + ["Status", "Zakres dokumentu", "Plik w folderze", "Pozostałe pliki",
                             "Data pobrania (UTC)", "Źródło – strona producenta (tekst)", "MPN potwierdzony w treści",
                             "Powód niepowodzenia / uwagi", "Do ręcznej weryfikacji (oficjalne strony)"]
COMPLIANCE_W = {**IDENT_W, "Status": 30, "Zakres dokumentu": 18, "Plik w folderze": 45, "Pozostałe pliki": 40,
                "Data pobrania (UTC)": 20, "Źródło – strona producenta (tekst)": 55,
                "MPN potwierdzony w treści": 12, "Powód niepowodzenia / uwagi": 60,
                "Do ręcznej weryfikacji (oficjalne strony)": 50}

LC_FILL = {LifecycleStatus.ACTIVE: GREEN, LifecycleStatus.MATURE: YELLOW, LifecycleStatus.PREVIEW: YELLOW,
           LifecycleStatus.NRND: ORANGE, LifecycleStatus.LAST_TIME_BUY: RED, LifecycleStatus.OBSOLETE: RED,
           LifecycleStatus.UNKNOWN: GREY}
LIFECYCLE_HEADER = IDENT + ["Status", "Etykieta producenta (dosłownie)", "Zakres statusu", "Sprawdzono (UTC)",
                            "Kopia strony w folderze", "Źródło – strona producenta (tekst)",
                            "Dowód (fragment strony)", "Uwagi"]
LIFECYCLE_W = {**IDENT_W, "Status": 30, "Etykieta producenta (dosłownie)": 28, "Zakres statusu": 20,
               "Sprawdzono (UTC)": 20, "Kopia strony w folderze": 40, "Źródło – strona producenta (tekst)": 50,
               "Dowód (fragment strony)": 60, "Uwagi": 50}


def _lifecycle_rows(results: list[ItemResult], lifecycle_pl: dict) -> list[list]:
    rows = []
    for n, r in enumerate(results, 1):
        lc = r.lifecycle
        if lc is None:
            rows.append([*_ident(n, r), Link("NIE SPRAWDZANO", None, GREY)] + [""] * 7)
            continue
        scope = {"part": "dla MPN", "page": "strona produktu / rodziny"}.get(lc.scope, "")
        rows.append([
            *_ident(n, r),
            Link(lifecycle_pl[lc.status], lc.snapshot or None, LC_FILL[lc.status]),
            lc.label, scope, lc.checked_at,
            Link(Path(lc.snapshot).name, lc.snapshot) if lc.snapshot else "",
            lc.source_url, lc.evidence, lc.note,
        ])
    return rows


LONGEVITY_HEADER = IDENT + ["Status", "Produkcja do (rok)", "Okres (lata)", "Od (rok)", "Podstawa daty",
                            "Zakres deklaracji", "Program / dokument", "Plik w folderze",
                            "Źródło – strona producenta (tekst)", "Dowód (fragment)", "Uwagi"]
LONGEVITY_W = {**IDENT_W, "Status": 34, "Produkcja do (rok)": 12, "Okres (lata)": 10, "Od (rok)": 10,
               "Podstawa daty": 26, "Zakres deklaracji": 20, "Program / dokument": 30, "Plik w folderze": 45,
               "Źródło – strona producenta (tekst)": 50, "Dowód (fragment)": 60, "Uwagi": 50}


def longevity_status(r: ItemResult) -> tuple[str, str]:
    lg = r.longevity
    if lg is None:
        return "NIE SPRAWDZANO", GREY
    if lg.found and lg.end_year:
        suffix = "" if lg.scope == "part" else " (deklaracja dla rodziny)"
        return f"DEKLARACJA: produkcja do {lg.end_year}{suffix}", GREEN if lg.scope == "part" else YELLOW
    if lg.found:
        return "OBJĘTY PROGRAMEM LONGEVITY – brak daty końcowej", YELLOW
    if lg.docs or lg.source_url:
        return "TYLKO OGÓLNA POLITYKA PRODUCENTA – brak deklaracji dla MPN", ORANGE
    return "BRAK DEKLARACJI – do uzyskania od producenta", RED


def _longevity_rows(results: list[ItemResult]) -> list[list]:
    rows = []
    for n, r in enumerate(results, 1):
        lg = r.longevity
        text, fill = longevity_status(r)
        if lg is None:
            rows.append([*_ident(n, r), Link(text, None, fill)] + [""] * 10)
            continue
        target = lg.snapshot or (lg.docs[0].path_for(DocType.LONGEVITY) if lg.docs else None)
        files = [Link(Path(target).name, target)] if target else [""]
        scope = {"part": "dla MPN", "family": "rodzina (prefiks MPN)", "general": "ogólna polityka"}.get(lg.scope, lg.scope)
        rows.append([
            *_ident(n, r), Link(text, target, fill),
            lg.end_year or "", lg.years or "", lg.start_year or "", lg.end_basis, scope, lg.program,
            files[0], lg.source_url, lg.evidence, lg.note,
        ])
    return rows


# ------------------------------------------------------------------ zapis

def write_workbook(path: Path, out_dir: Path, results: list[ItemResult], summary_rows: list[tuple],
                   meta: dict, settings, item_header: list[str], items: list[list], file_header: list[str],
                   files: list[list], groups, req_rows: list[list], req_header: list[str], invalid,
                   status_pl: dict, lifecycle_pl: dict) -> Path:
    w = _Writer(out_dir)

    # Podsumowanie
    ws = w.wb.active
    ws.title = "Podsumowanie"
    ws.append(["Raport RoHS / REACH / cykl życia / długość produkcji"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Wygenerowano (UTC)", datetime.now(timezone.utc).replace(microsecond=0).isoformat()])
    ws.append(["Arkusz / nagłówek BoM", f"{meta.get('sheet')} / wiersz {meta.get('header_row')}"])
    ws.append(["Mapowanie kolumn", ", ".join(f"{k}={v}" for k, v in (meta.get("columns") or {}).items())
               + (f" (wykryte na podstawie {meta.get('detection')})" if meta.get("detection") else "")])
    for line in meta.get("sheets", []):
        ws.append(["Analiza arkusza", line])
    for warn in meta.get("warnings", []):
        ws.append(["UWAGA (BoM)", warn])
    ws.append([])
    ws.append(["Miara", "Liczba", "Procent pozycji"])
    for c in ws[ws.max_row]:
        c.font, c.fill = w.bold, w.head_fill
    for row in summary_rows:
        ws.append(list(row))
    ws.append([])
    ws.append(["Foldery z pobranymi plikami", "(kliknij, aby otworzyć)"])
    ws[ws.max_row][0].font = w.bold
    docs_dir = out_dir / "documents"
    for label, folder in (("RoHS", TYPE_FOLDERS[DocType.ROHS]), ("REACH", TYPE_FOLDERS[DocType.REACH]),
                          ("Status cyklu życia (kopie stron)", LIFECYCLE_FOLDER),
                          ("Długość produkcji", TYPE_FOLDERS[DocType.LONGEVITY]),
                          ("Deklaracje materiałowe (MCD)", TYPE_FOLDERS[DocType.MCD])):
        target = docs_dir / folder
        ws.append([label, f"documents/{folder}" if target.is_dir() else f"documents/{folder} (brak plików)"])
        if target.is_dir():
            cell = ws.cell(row=ws.max_row, column=2)
            cell.hyperlink = w.rel(str(target))
            cell.font = w.link_font
    ws.append([])
    ws.append(["Zasady", "Wszystkie pliki pochodzą wyłącznie z oficjalnych domen producentów. Kolumna 'Status' "
                         "na kartach RoHS / REACH / Status cyklu życia / Długość produkcji jest linkiem do pliku "
                         "w folderze 'documents' (link względny – przenoś cały folder raportu). Adres strony "
                         "producenta podany jest jako tekst w kolumnie 'Źródło'."])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 16
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # Karty per temat
    w.sheet("RoHS", COMPLIANCE_HEADER, _compliance_rows(results, DocType.ROHS, status_pl), COMPLIANCE_W)
    w.sheet("REACH", COMPLIANCE_HEADER, _compliance_rows(results, DocType.REACH, status_pl), COMPLIANCE_W)
    w.sheet("Status cyklu życia", LIFECYCLE_HEADER, _lifecycle_rows(results, lifecycle_pl), LIFECYCLE_W)
    w.sheet("Długość produkcji", LONGEVITY_HEADER, _longevity_rows(results), LONGEVITY_W)

    # Szczegóły (wszystkie kolumny) – statusy RoHS/REACH również jako linki do plików
    det_rows = []
    i_rohs, i_reach = item_header.index("Status RoHS"), item_header.index("Status REACH")
    for r, row in zip(results, items):
        row = list(row)
        for idx, t, st in ((i_rohs, DocType.ROHS, r.rohs_status), (i_reach, DocType.REACH, r.reach_status)):
            docs = r.docs_for(t)
            row[idx] = Link(row[idx], docs[0].path_for(t) if docs else None, _status_fill(st))
        det_rows.append(row)
    w.sheet("Szczegóły pozycji", item_header, det_rows,
            {"Plik RoHS": 50, "URL źródłowy RoHS": 55, "Plik REACH": 50, "URL źródłowy REACH": 55,
             "Powód niepowodzenia / uwagi": 60, "Uwagi do MPN": 50, "Status RoHS": 28, "Status REACH": 28})

    file_rows_linked = [[Link(row[0], row[0]), *row[1:]] for row in files]
    w.sheet("Pliki", file_header, file_rows_linked, {"Ścieżka": 70, "URL źródłowy": 60, "URL końcowy": 60})
    w.sheet("Do uzyskania mailowo", req_header, req_rows,
            {"Producent": 22, "MPN": 30, "Do uzyskania": 40, "Status": 50, "Kontakt (oficjalna strona producenta)": 90})
    w.sheet("Szablony e-mail", ["Producent", "Liczba MPN", "Treść e-maila"],
            [[g.manufacturer, len(g.items), g.email] for g in groups],
            {"Producent": 22, "Liczba MPN": 10, "Treść e-maila": 120})
    w.sheet("Niepoprawne wiersze", ["Wiersz BoM", "Powód", "Zawartość"],
            [[r.row_number, r.reason, "; ".join(f"{k}={v}" for k, v in r.raw.items())] for r in invalid],
            {"Wiersz BoM": 12, "Powód": 50, "Zawartość": 100})
    w.wb.save(path)
    return path
