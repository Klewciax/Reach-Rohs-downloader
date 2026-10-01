"""Raport końcowy: CSV + XLSX + szablony e-maili + podsumowanie w konsoli."""
from __future__ import annotations

import csv
import logging
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from .config import Settings
from .contacts import ContactInfo
from .downloader import safe_name
from .models import DocType, InvalidRow, ItemResult, Status
from .pipeline import is_success

log = logging.getLogger(__name__)

STATUS_PL = {
    Status.DOWNLOADED: "POBRANO (dokument dla MPN)",
    Status.DOWNLOADED_FAMILY: "POBRANO – DOKUMENT ZBIORCZY (rodzina)",
    Status.GENERAL_ONLY: "TYLKO OGÓLNE OŚWIADCZENIE PRODUCENTA",
    Status.LOGIN_REQUIRED: "WYMAGA LOGOWANIA",
    Status.FORM_REQUIRED: "DOSTĘPNE PRZEZ FORMULARZ",
    Status.BLOCKED_ROBOTS: "ZABLOKOWANE (robots.txt)",
    Status.NOT_FOUND: "NIE ZNALEZIONO",
    Status.ERROR: "BŁĄD",
    Status.UNKNOWN_MANUFACTURER: "NIEZNANY PRODUCENT",
}

ITEM_COLUMNS = [
    "MPN", "Manufacturer (BoM)", "Manufacturer (znormalizowany)", "Dopasowanie producenta", "Ref. designators",
    "Wiersz(e) w pliku BoM", "Status RoHS", "Status REACH", "Plik RoHS", "URL źródłowy RoHS", "Zakres RoHS",
    "Plik REACH", "URL źródłowy REACH", "Zakres REACH", "Data pobrania (UTC)", "MPN potwierdzony w treści",
    "Powód niepowodzenia / uwagi", "Do ręcznego sprawdzenia (oficjalne strony)",
]


@dataclass
class Summary:
    bom_rows_total: int
    bom_rows_valid: int
    bom_rows_invalid: int
    items: int
    rohs: int
    reach: int
    both: int
    none: int
    family_docs: int
    extra: dict = field(default_factory=dict)

    def pct(self, n: int) -> str:
        return f"{(100.0 * n / self.items):.1f}%" if self.items else "0.0%"


def summarize(results: list[ItemResult], settings: Settings, meta: dict, invalid: list[InvalidRow]) -> Summary:
    rohs = sum(is_success(r.rohs_status, settings) for r in results)
    reach = sum(is_success(r.reach_status, settings) for r in results)
    both = sum(is_success(r.rohs_status, settings) and is_success(r.reach_status, settings) for r in results)
    none = sum(not is_success(r.rohs_status, settings) and not is_success(r.reach_status, settings) for r in results)
    fam = sum(Status.DOWNLOADED_FAMILY in (r.rohs_status, r.reach_status) for r in results)
    return Summary(meta.get("total_rows", 0), meta.get("total_rows", 0) - len(invalid), len(invalid),
                   len(results), rohs, reach, both, none, fam)


def _doc_cells(r: ItemResult, t: DocType) -> tuple[str, str, str, str]:
    docs = r.docs_for(t)
    if docs:  # pokazuj tylko dokumenty o najlepszym (najwęższym) zakresie
        docs = [d for d in docs if d.scope == docs[0].scope]
    return ("\n".join(d.path for d in docs), "\n".join(d.final_url if d.final_url == d.url else
                                                       f"{d.url} -> {d.final_url}" for d in docs),
            "\n".join(d.scope.value for d in docs), "\n".join(dict.fromkeys(d.downloaded_at for d in docs)))


def item_rows(results: list[ItemResult]) -> list[list[str]]:
    rows = []
    for r in results:
        rp, ru, rs, rd = _doc_cells(r, DocType.ROHS)
        ep, eu, es, ed = _doc_cells(r, DocType.REACH)
        dates = "\n".join(dict.fromkeys(x for x in (rd + "\n" + ed).split("\n") if x))
        verified = "; ".join(dict.fromkeys(f"{Path(d.path).name}: {d.mpn_verified}" for d in r.docs))
        reasons = r.reasons + [n for n in r.notes if n not in r.reasons]
        for d in r.docs:
            if d.note:
                reasons.append(f"{Path(d.path).name}: {d.note}")
        rows.append([
            r.item.mpn, " | ".join(r.item.manufacturer_variants) or r.item.manufacturer_raw,
            r.item.manufacturer.name if r.item.manufacturer else "", r.item.match_method,
            ", ".join(r.item.refdes), ", ".join(map(str, r.item.rows)),
            STATUS_PL[r.rohs_status], STATUS_PL[r.reach_status], rp, ru, rs, ep, eu, es, dates, verified,
            "\n".join(reasons), "\n".join(r.manual_urls),
        ])
    return rows


def file_rows(results: list[ItemResult]) -> list[list[str]]:
    seen: "OrderedDict[str, list]" = OrderedDict()
    for r in results:
        for d in r.docs:
            if d.path not in seen:
                seen[d.path] = [d.path, d.url, d.final_url, d.downloaded_at, d.sha256,
                                ", ".join(sorted(t.value for t in d.doc_types)), d.scope.value,
                                r.item.manufacturer_name, [], d.title]
            seen[d.path][8].append(r.item.mpn)
    return [[*v[:8], ", ".join(v[8]), v[9]] for v in seen.values()]


@dataclass
class RequestGroup:
    manufacturer: str
    items: list[tuple[str, list[tuple[str, bool]], str]]  # (MPN, brakujące typy, status)
    contact: ContactInfo | None
    email: str = ""


def missing_types(r: ItemResult, settings: Settings) -> list[tuple[str, bool]]:
    """Lista (rodzaj, czy_jest_tylko_dokument_zbiorczy) dla dokumentów do uzyskania od producenta."""
    out = []
    for name, st in (("RoHS", r.rohs_status), ("REACH", r.reach_status)):
        if st == Status.DOWNLOADED_FAMILY:
            out.append((name, True))  # prosimy o deklarację dla konkretnego MPN
        elif not is_success(st, settings):
            out.append((name, False))
    return out


def _missing_pl(miss: list[tuple[str, bool]]) -> str:
    return ", ".join(f"{t} (jest tylko dokument zbiorczy – potrzebna deklaracja dla MPN)" if fam else t
                     for t, fam in miss)


def _missing_en(miss: list[tuple[str, bool]]) -> str:
    return ", ".join(f"{t} – part-specific (only a product-family document found)" if fam else t
                     for t, fam in miss)


def build_request_groups(results: list[ItemResult], settings: Settings,
                         contacts: dict[str, ContactInfo]) -> list[RequestGroup]:
    groups: "OrderedDict[str, RequestGroup]" = OrderedDict()
    for r in results:
        miss = missing_types(r, settings)
        if not miss:
            continue
        name = r.item.manufacturer_name
        key = r.item.manufacturer.key if r.item.manufacturer else "?" + name
        g = groups.get(key)
        if g is None:
            g = groups[key] = RequestGroup(name, [], contacts.get(key))
        status = f"RoHS: {STATUS_PL[r.rohs_status]}; REACH: {STATUS_PL[r.reach_status]}"
        g.items.append((r.item.mpn, miss, status))
    for g in groups.values():
        g.email = email_template(g, settings)
    return list(groups.values())


def email_template(g: RequestGroup, s: Settings) -> str:
    lines = [f"  - {mpn}  (needed: {_missing_en(m)})" for mpn, m, _ in g.items]
    return f"""Subject: Request for RoHS and REACH (SVHC) compliance declarations – {len(g.items)} {g.manufacturer} part number(s)

Dear {g.manufacturer} Product Compliance / Environmental Team,

We use the following {g.manufacturer} components in our product "{s.project_name}" and are
compiling the environmental compliance documentation for it. We were unable to obtain
part-specific declarations for these part numbers from your official website:

{chr(10).join(lines)}

For each part number listed above, could you please provide:

  1. EU RoHS declaration / certificate of compliance (Directive 2011/65/EU as amended by
     (EU) 2015/863), stating any RoHS exemptions used (Annex III/IV item numbers).
  2. EU REACH declaration (Regulation (EC) No 1907/2006), including the SVHC status
     against the latest ECHA Candidate List: name and CAS number of any SVHC present
     above 0.1% w/w, its concentration and location in the article, and the Candidate
     List version the statement refers to.
  3. If available, a full material declaration (IPC-1752A Class D or equivalent).

The documents should be issued by {g.manufacturer} as the manufacturer and clearly reference
the exact manufacturer part numbers (orderable part numbers) listed above. A signed PDF
is preferred.

Thank you in advance for your support.

Best regards,
{s.requester_name}
{s.requester_company}
{s.requester_email}
"""


def write_reports(out_dir: Path, results: list[ItemResult], invalid: list[InvalidRow], meta: dict,
                  settings: Settings, groups: list[RequestGroup], summary: Summary) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    items = item_rows(results)
    files = file_rows(results)

    def write_csv(name: str, header: list[str], rows: list[list]) -> Path:
        p = out_dir / name
        with open(p, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(header)
            w.writerows(rows)
        return p

    paths["items_csv"] = write_csv("report_items.csv", ITEM_COLUMNS, items)
    file_header = ["Ścieżka", "URL źródłowy", "URL końcowy", "Data pobrania (UTC)", "SHA-256", "Rodzaj",
                   "Zakres", "Producent", "Użyty dla MPN", "Tytuł"]
    paths["files_csv"] = write_csv("report_files.csv", file_header, files)
    req_header = ["Producent", "MPN", "Do uzyskania", "Status", "Kontakt (oficjalna strona producenta)"]
    req_rows = []
    for g in groups:
        contact = g.contact.summary() if g.contact else \
            "DO RĘCZNEJ WERYFIKACJI: producent spoza rejestru – brak zweryfikowanej domeny"
        for mpn, miss, status in g.items:
            req_rows.append([g.manufacturer, mpn, _missing_pl(miss), status, contact])
    paths["requests_csv"] = write_csv("report_to_request.csv", req_header, req_rows)

    tpl_dir = out_dir / "email_templates"
    tpl_dir.mkdir(exist_ok=True)
    for g in groups:
        p = tpl_dir / f"{safe_name(g.manufacturer)}.txt"
        contact = g.contact.summary() if g.contact else "DO RĘCZNEJ WERYFIKACJI"
        p.write_text(f"# Kontakt:\n# " + contact.replace("\n", "\n# ") + "\n\n" + g.email, encoding="utf-8")

    paths["xlsx"] = _write_xlsx(out_dir / "report.xlsx", summary, items, files, file_header, groups, req_rows,
                                req_header, invalid, meta, settings)
    return paths


def summary_lines(s: Summary) -> list[tuple[str, str, str]]:
    return [
        ("Wiersze BoM (niepuste)", str(s.bom_rows_total), ""),
        ("Wiersze poprawne", str(s.bom_rows_valid), ""),
        ("Wiersze niepoprawne / pominięte", str(s.bom_rows_invalid), ""),
        ("Pozycje BoM po deduplikacji (producent + MPN)", str(s.items), "100%"),
        ("Pobrano RoHS", str(s.rohs), s.pct(s.rohs)),
        ("Pobrano REACH", str(s.reach), s.pct(s.reach)),
        ("Pobrano oba (RoHS i REACH)", str(s.both), s.pct(s.both)),
        ("Nie pobrano żadnego", str(s.none), s.pct(s.none)),
        ("  w tym pozycje z dokumentem zbiorczym (rodzina)", str(s.family_docs), s.pct(s.family_docs)),
    ]


def _write_xlsx(path: Path, s: Summary, items, files, file_header, groups, req_rows, req_header,
                invalid: list[InvalidRow], meta: dict, settings: Settings) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    fills = {
        "POBRANO (": PatternFill("solid", fgColor="C6EFCE"),
        "POBRANO –": PatternFill("solid", fgColor="FFF2CC"),
        "TYLKO": PatternFill("solid", fgColor="FFF2CC"),
    }
    bad = PatternFill("solid", fgColor="FFC7CE")

    def sheet(ws, header, rows, widths=None, status_cols=()):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head_fill
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, _ in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 22)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for col in status_cols:
                cell = row[col - 1]
                val = str(cell.value or "")
                cell.fill = next((f for k, f in fills.items() if val.startswith(k)), bad)

    ws = wb.active
    ws.title = "Podsumowanie"
    ws.append(["Raport zgodności RoHS / REACH", "", ""])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Wygenerowano (UTC)", datetime.now(timezone.utc).replace(microsecond=0).isoformat(), ""])
    ws.append(["Arkusz / nagłówek BoM", f"{meta.get('sheet')} / wiersz {meta.get('header_row')}", ""])
    ws.append(["Mapowanie kolumn", ", ".join(f"{k}={v}" for k, v in (meta.get('columns') or {}).items()), ""])
    ws.append([])
    ws.append(["Miara", "Liczba", "Procent pozycji"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, head_fill
    for row in summary_lines(s):
        ws.append(list(row))
    ws.append([])
    ws.append(["Zasady liczenia", (
        "Sukces = dokument dla konkretnego MPN" +
        (" lub dokument zbiorczy dla rodziny" if settings.count_family_as_success else "") +
        (" lub ogólne oświadczenie producenta" if settings.count_general_as_success else "") +
        ". Wszystkie pliki pochodzą wyłącznie z oficjalnych domen producentów."), ""])
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 16

    widths = {1: 22, 2: 26, 3: 22, 9: 50, 10: 60, 12: 50, 13: 60, 17: 70, 18: 60}
    sheet(wb.create_sheet("Pozycje"), ITEM_COLUMNS, items, widths, status_cols=(7, 8))
    sheet(wb.create_sheet("Pliki"), file_header, files, {1: 70, 2: 70, 3: 70, 5: 20})
    sheet(wb.create_sheet("Do uzyskania mailowo"), req_header, req_rows, {1: 24, 2: 24, 3: 16, 4: 50, 5: 90})
    sheet(wb.create_sheet("Szablony e-mail"), ["Producent", "Liczba MPN", "Treść e-maila"],
          [[g.manufacturer, len(g.items), g.email] for g in groups], {1: 24, 2: 12, 3: 120})
    sheet(wb.create_sheet("Niepoprawne wiersze"), ["Wiersz BoM", "Powód", "Zawartość"],
          [[r.row_number, r.reason, "; ".join(f"{k}={v}" for k, v in r.raw.items())] for r in invalid],
          {1: 12, 2: 50, 3: 100})
    wb.save(path)
    return path


def print_console_summary(s: Summary, groups: list[RequestGroup], paths: dict[str, Path], results: list[ItemResult]
                          ) -> None:
    w = 52
    print("\n" + "=" * 72)
    print("RAPORT RoHS / REACH")
    print("=" * 72)
    for label, n, pct in summary_lines(s):
        print(f"{label:<{w}} {n:>6}  {pct:>7}")
    status_counts: dict[str, int] = {}
    for r in results:
        for st in (r.rohs_status, r.reach_status):
            status_counts[st.value] = status_counts.get(st.value, 0) + 1
    print("-" * 72)
    print("Statusy (RoHS + REACH łącznie): " + ", ".join(f"{k}={v}" for k, v in sorted(status_counts.items())))
    if groups:
        print("-" * 72)
        print("Do uzyskania mailowo:")
        for g in groups:
            print(f"  {g.manufacturer}: {len(g.items)} MPN – " + ", ".join(m for m, _, _ in g.items[:8])
                  + (" …" if len(g.items) > 8 else ""))
    print("-" * 72)
    for k, p in paths.items():
        print(f"{k:<14} {p}")
    print("=" * 72)
