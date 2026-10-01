"""Wczytywanie BoM (CSV/XLSX), mapowanie kolumn, walidacja i deduplikacja."""
from __future__ import annotations

import csv
import logging
import re
from pathlib import Path

from .classify import compact
from .manufacturers import ManufacturerRegistry, is_generic_manufacturer, normalize_name, strip_legal
from .models import BomItem, BomRow, InvalidRow

log = logging.getLogger(__name__)

# Typowe nazwy kolumn (porównywane po normalizacji: małe litery, bez znaków niealfanumerycznych).
COLUMN_SYNONYMS: dict[str, list[str]] = {
    "mpn": [
        "mpn", "manufacturer part number", "manufacturer part no", "manufacturer pn",
        "mfr part number", "mfr part no", "mfr part", "mfr pn", "mfr p n", "mfg part number",
        "mfg pn", "mfg part no", "manufacturer part", "mfr number", "part number manufacturer",
        "mpn1", "manufacturer part number 1", "numer katalogowy producenta", "nr kat producenta",
        "part number", "partnumber", "pn", "p n",
    ],
    "manufacturer": [
        "manufacturer", "manufacturer name", "mfr", "mfr name", "mfg", "mfg name", "maker",
        "manufacturer 1", "manufacturer1", "mfr1", "producent", "brand",
    ],
    "refdes": [
        "refdes", "ref des", "reference", "references", "reference designator",
        "reference designators", "designator", "designators", "ref", "refs", "oznaczenie",
        "part reference", "pozycja",
    ],
    "quantity": ["qty", "quantity", "ilosc", "ilość", "count", "qty per board"],
}

INVALID_MPN_VALUES = {"", "NA", "N/A", "-", "--", "TBD", "DNP", "DNF", "NC", "NONE", "NOTUSED", "?", "X"}


def _norm_header(h: object) -> str:
    return re.sub(r"[^a-z0-9ąćęłńóśźż]+", " ", str(h or "").lower()).strip()


def _detect_columns(headers: list[str], explicit: dict[str, str]) -> dict[str, int]:
    norm = [_norm_header(h) for h in headers]
    mapping: dict[str, int] = {}
    for logical, colname in (explicit or {}).items():
        if not colname:
            continue
        target = _norm_header(colname)
        if target not in norm:
            raise ValueError(
                f"Kolumna '{colname}' (dla '{logical}') nie istnieje w BoM. Dostępne: {', '.join(map(str, headers))}"
            )
        mapping[logical] = norm.index(target)
    for logical, synonyms in COLUMN_SYNONYMS.items():
        if logical in mapping:
            continue
        for syn in synonyms:  # kolejność synonimów = priorytet
            if syn in norm and norm.index(syn) not in mapping.values():
                mapping[logical] = norm.index(syn)
                break
    return mapping


def _read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect)]


def _read_xlsx(path: Path, sheet: str | None) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = [wb[sheet]] if sheet else wb.worksheets
    out = []
    for ws in sheets:
        rows = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
        out.append((ws.title, rows))
    return out


def _find_header(rows: list[list[str]], explicit: dict[str, str]) -> tuple[int, dict[str, int]] | None:
    for i, row in enumerate(rows[:30]):
        try:
            mapping = _detect_columns(row, explicit)
        except ValueError:
            continue
        if "mpn" in mapping and "manufacturer" in mapping:
            return i, mapping
    return None


def read_bom(path: str | Path, columns: dict[str, str] | None = None, sheet: str | None = None
             ) -> tuple[list[BomRow], list[InvalidRow], dict]:
    """Zwraca (poprawne wiersze, niepoprawne wiersze, metadane)."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Nie znaleziono pliku BoM: {path}")
    columns = columns or {}
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        tables = _read_xlsx(path, sheet)
    elif suffix in (".csv", ".txt", ".tsv"):
        tables = [("csv", _read_csv(path))]
    else:
        raise ValueError(f"Nieobsługiwany format BoM: {suffix} (obsługiwane: .csv, .tsv, .xlsx)")

    found = None
    for title, rows in tables:
        hdr = _find_header(rows, columns)
        if hdr:
            found = (title, rows, *hdr)
            break
    if not found:
        # Powtórz z jawnym mapowaniem, by zwrócić czytelny błąd.
        if columns and tables:
            _detect_columns(tables[0][1][0] if tables[0][1] else [], columns)
        raise ValueError(
            "Nie znaleziono nagłówka z kolumnami MPN i Manufacturer. "
            "Podaj mapowanie, np. --col mpn=\"Mfr Part #\" --col manufacturer=\"Mfr\""
        )
    title, rows, header_idx, mapping = found
    headers = rows[header_idx]
    log.info("BoM: arkusz '%s', nagłówek w wierszu %d, mapowanie: %s", title, header_idx + 1,
             {k: headers[v] for k, v in mapping.items()})

    def cell(row: list[str], key: str) -> str:
        idx = mapping.get(key)
        if idx is None or idx >= len(row):
            return ""
        return str(row[idx] or "").strip()

    valid: list[BomRow] = []
    invalid: list[InvalidRow] = []
    empty = 0
    for i, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(str(c).strip() for c in row):
            empty += 1
            continue
        raw = {str(headers[j]): row[j] for j in range(min(len(headers), len(row))) if str(row[j]).strip()}
        mpn = re.sub(r"\s+", " ", cell(row, "mpn"))
        mfr = re.sub(r"\s+", " ", cell(row, "manufacturer"))
        if compact(mpn) == "" or mpn.upper().replace(" ", "") in INVALID_MPN_VALUES:
            invalid.append(InvalidRow(i, "Brak lub niepoprawny MPN", raw))
            continue
        if not mfr or is_generic_manufacturer(mfr):
            invalid.append(InvalidRow(i, f"Brak konkretnego producenta (MPN {mpn})", raw))
            continue
        if len(mpn) > 64:
            invalid.append(InvalidRow(i, "MPN podejrzanie długi (>64 znaki) – prawdopodobnie opis, nie MPN", raw))
            continue
        refdes = [r for r in re.split(r"[,;\s]+", cell(row, "refdes")) if r]
        valid.append(BomRow(i, mpn, mfr, refdes, cell(row, "quantity")))
    meta = {"sheet": title, "header_row": header_idx + 1,
            "columns": {k: headers[v] for k, v in mapping.items()},
            "total_rows": len(valid) + len(invalid), "empty_rows": empty}
    return valid, invalid, meta


def group_items(rows: list[BomRow], registry: ManufacturerRegistry) -> list[BomItem]:
    """Deduplikacja po (producent, MPN). Ten sam MPN u różnych producentów = różne pozycje."""
    items: dict[tuple[str, str], BomItem] = {}
    for r in rows:
        mfr, method = registry.resolve(r.manufacturer)
        mkey = mfr.key if mfr else "?" + strip_legal(normalize_name(r.manufacturer))
        key = (mkey, compact(r.mpn))
        item = items.get(key)
        if item is None:
            item = BomItem(mpn=r.mpn, manufacturer_raw=r.manufacturer, manufacturer=mfr, match_method=method)
            items[key] = item
        item.rows.append(r.row_number)
        for d in r.refdes:
            if d not in item.refdes:
                item.refdes.append(d)
        if r.manufacturer not in item.manufacturer_variants:
            item.manufacturer_variants.append(r.manufacturer)
    return list(items.values())
